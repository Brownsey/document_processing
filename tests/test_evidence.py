import base64
import json
import os
import subprocess
from decimal import Decimal
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document

from agent_pipeline.contracts import Err, ModelReply, Ok, ProviderError
from agent_pipeline.evidence import load_sources

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+a3ioAAAAASUVORK5CYII="
)


class ImageReader:
    def __init__(self, reply=None):
        self.calls = []
        self.reply = reply

    def complete(self, **kwargs):
        self.calls.append(kwargs)
        assert kwargs["image_path"].is_file()
        assert "holders" not in json.dumps(kwargs["context"])
        return self.reply or Ok(
            ModelReply(
                text="",
                data={
                    "text": "Only on image: account Q1 GBP 123 on 2026-03-01",
                    "complete": True,
                },
                usage={},
                model="fake",
            )
        )


def test_docx_preserves_paragraph_table_order_and_stable_references(tmp_path):
    doc = Document()
    doc.add_paragraph("Before table")
    row = doc.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text = "Account Q1"
    row.cells[1].text = "GBP 123"
    doc.add_paragraph("After table")
    doc.save(tmp_path / "renamed.docx")
    first = load_sources(tmp_path, ImageReader())
    second = load_sources(tmp_path, ImageReader())
    assert isinstance(first, Ok)
    assert isinstance(second, Ok)
    assert [b.text for b in first.value.blocks] == [
        "Before table",
        "Account Q1 | GBP 123",
        "After table",
    ]
    assert [b.locator for b in first.value.blocks] == [
        "body/paragraph[1]",
        "body/table[1]/row[1]",
        "body/paragraph[2]",
    ]
    assert first.value.blocks == second.value.blocks
    assert len({b.id for b in first.value.blocks}) == 3
    assert all(len(b.sha256) == 64 for b in first.value.blocks)


def test_nested_renamed_guidance_and_marketing_are_separated_without_losing_novel_text(
    tmp_path,
):
    folder = tmp_path / "nested"
    folder.mkdir()
    (folder / "one.md").write_text(
        "# Internal notes: data sources\nFor whoever configures the report: confirm fee rates.",
        encoding="utf-8",
    )
    (folder / "two.txt").write_text(
        "Quarterly update\nThis note is general market commentary and does not relate to any "
        "individual client's accounts or recommendations. Illustrative portfolio GBP 900.",
        encoding="utf-8",
    )
    (folder / "three.md").write_text(
        "# New receipt\nQ1 received GBP 234.", encoding="utf-8"
    )
    (folder / "four.md").write_text(
        "# Report specification: Advice\nWhat each section of the report must contain.",
        encoding="utf-8",
    )
    result = load_sources(tmp_path, ImageReader())
    assert isinstance(result, Ok)
    roles = {b.source: b.role for b in result.value.blocks}
    assert roles == {
        "nested/one.md": "guidance",
        "nested/two.txt": "excluded",
        "nested/three.md": "evidence",
        "nested/four.md": "guidance",
    }
    assert any("GBP 234" in b.text for b in result.value.blocks)


def test_json_objects_remain_structured_and_unique_image_is_read_independently(
    tmp_path,
):
    raw = {"holders": {"client": {"accounts": [{"account_id": "Q1", "value": 50}]}}}
    (tmp_path / "renamed.json").write_text(json.dumps(raw), encoding="utf-8")
    (tmp_path / "unique.png").write_bytes(PNG)
    reader = ImageReader()
    result = load_sources(tmp_path, reader)
    assert isinstance(result, Ok)
    assert result.value.databases == [raw]
    assert any(
        "GBP 123" in b.text and b.source == "unique.png" for b in result.value.blocks
    )
    assert len(reader.calls) == 1
    assert reader.calls[0]["task"] == "read_image"
    assert all(i["status"] == "read" for i in result.value.inventory)


@pytest.mark.parametrize("content", ["{broken", "42", "[1,2]"])
def test_malformed_json_fails_before_any_image_call(tmp_path, content):
    (tmp_path / "a.png").write_bytes(PNG)
    (tmp_path / "z.json").write_text(content, encoding="utf-8")
    reader = ImageReader()
    result = load_sources(tmp_path, reader)
    assert isinstance(result, Err)
    assert result.error.code == "invalid_source"
    assert reader.calls == []


def test_json_list_records_are_preserved(tmp_path):
    (tmp_path / "new.json").write_text('[{"account_id":"X"},{"account_id":"Y"}]')
    result = load_sources(tmp_path, ImageReader())
    assert isinstance(result, Ok)
    assert result.value.databases == [{"account_id": "X"}, {"account_id": "Y"}]


def test_unknown_file_is_material_unresolved_inventory(tmp_path):
    (tmp_path / "new.pdf").write_bytes(b"%PDF unfamiliar source")
    result = load_sources(tmp_path, ImageReader())
    assert isinstance(result, Err)
    item = result.error.details["inventory"][0]
    assert item["source"] == "new.pdf"
    assert item["status"] == "unresolved"
    assert item["material"] is True
    assert item["issues"] == ["unsupported_file_type"]


def test_embedded_docx_content_is_not_silently_dropped(tmp_path):
    path = tmp_path / "embedded.docx"
    doc = Document()
    doc.add_paragraph("Text outside image")
    doc.save(path)
    with ZipFile(path) as archive:
        files = {n: archive.read(n) for n in archive.namelist()}
    files["word/document.xml"] = files["word/document.xml"].replace(
        b"</w:body>", b"<w:p><w:r><w:drawing/></w:r></w:p></w:body>"
    )
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, value in files.items():
            archive.writestr(name, value)
    result = load_sources(tmp_path, ImageReader())
    assert isinstance(result, Err)
    inventory = result.error.details["inventory"]
    assert inventory[0]["status"] == "unresolved"
    assert "unsupported_embedded_content" in inventory[0]["issues"]
    assert "Text outside image" not in str(result.error)


@pytest.mark.parametrize("kind", ["unsupported", "empty", "embedded"])
def test_material_unresolved_source_blocks_paid_image_calls_during_preflight(
    tmp_path, kind
):
    (tmp_path / "a.png").write_bytes(PNG)
    if kind == "unsupported":
        (tmp_path / "z.pdf").write_bytes(b"%PDF private source")
    elif kind == "empty":
        (tmp_path / "z.txt").write_text("")
    else:
        document = Document()
        paragraph = document.add_paragraph("Private client content")
        document.add_comment(paragraph.runs, text="Do not proceed", author="Adviser")
        document.save(tmp_path / "z.docx")
    reader = ImageReader()
    result = load_sources(tmp_path, reader)
    assert reader.calls == []
    assert isinstance(result, Err)
    assert result.error.code == "unreadable_evidence"
    inventory = result.error.details["inventory"]
    assert any(
        item["status"] == "unresolved" and item["material"] for item in inventory
    )
    assert "Private client content" not in str(result.error)
    assert "%PDF private source" not in str(result.error)


def test_image_provider_failure_is_preserved(tmp_path):
    (tmp_path / "unique.png").write_bytes(PNG)
    error = ProviderError("image_unavailable", "Image failed", stage="image")
    result = load_sources(tmp_path, ImageReader(Err(error)))
    assert isinstance(result, Err)
    assert result.error is error


@pytest.mark.parametrize(
    "data",
    [
        {"text": "", "complete": True},
        {"text": "Partial account", "complete": False},
        None,
    ],
)
def test_incomplete_image_transcription_is_failure(tmp_path, data):
    (tmp_path / "unique.png").write_bytes(PNG)
    reply = Ok(ModelReply(text="", data=data, usage={}, model="fake"))
    result = load_sources(tmp_path, ImageReader(reply))
    assert isinstance(result, Err)
    assert result.error.code == "incomplete_image"


@pytest.mark.parametrize(
    "name, content",
    [("bad.docx", b"invalid"), ("bad.png", b"invalid"), ("bad.txt", b"\xff\xfe\x00")],
)
def test_corrupt_sources_return_input_error(tmp_path, name, content):
    (tmp_path / name).write_bytes(content)
    assert isinstance(load_sources(tmp_path, ImageReader()), Err)


def test_size_limits_preflight_before_image_calls(tmp_path):
    (tmp_path / "a.png").write_bytes(PNG)
    with (tmp_path / "large.txt").open("wb") as file:
        file.truncate(25 * 1024 * 1024)
    reader = ImageReader()
    result = load_sources(tmp_path, reader)
    assert isinstance(result, Err)
    assert result.error.code == "source_too_large"
    assert reader.calls == []


def test_cross_client_symlink_is_rejected(tmp_path):
    client = tmp_path / "client"
    client.mkdir()
    other = tmp_path / "other.txt"
    other.write_text("Private other-client information")
    try:
        (client / "linked.txt").symlink_to(other)
    except OSError:
        outside = tmp_path / "outside"
        outside.mkdir()
        (outside / "private.txt").write_text("Private other-client information")
        subprocess.run(
            ["cmd", "/c", "mklink", "/J", str(client / "linked"), str(outside)],
            check=True,
            capture_output=True,
        )
    result = load_sources(client, ImageReader())
    assert isinstance(result, Err)
    assert result.error.code == "unsafe_source_path"


def test_cross_client_hardlink_is_rejected(tmp_path):
    client = tmp_path / "client"
    client.mkdir()
    other = tmp_path / "private.txt"
    other.write_text("Private other-client information")
    os.link(other, client / "linked.txt")
    result = load_sources(client, ImageReader())
    assert isinstance(result, Err)
    assert result.error.code == "unsafe_source_path"


def test_unreadable_subdirectory_fails_instead_of_silently_omitting_source(
    tmp_path, monkeypatch
):
    (tmp_path / "a.png").write_bytes(PNG)
    denied = tmp_path / "unreadable"
    denied.mkdir()
    original = os.scandir

    def permission_denied(path):
        if str(path).endswith("unreadable"):
            raise PermissionError("private detail must not enter errors")
        return original(path)

    monkeypatch.setattr(os, "scandir", permission_denied)
    reader = ImageReader()
    result = load_sources(tmp_path, reader)
    assert isinstance(result, Err)
    assert result.error.code == "unreadable_source"
    assert "private detail" not in str(result.error)
    assert reader.calls == []


def test_traversal_root_and_missing_directory_are_rejected(tmp_path):
    assert isinstance(load_sources(tmp_path / ".." / tmp_path.name, ImageReader()), Err)
    assert isinstance(load_sources(tmp_path / "missing", ImageReader()), Err)


def test_docx_archive_traversal_is_rejected(tmp_path):
    doc = Document()
    doc.add_paragraph("Safe text")
    path = tmp_path / "unsafe.docx"
    doc.save(path)
    with ZipFile(path, "a") as archive:
        archive.writestr("../outside.txt", "Untrusted")
    result = load_sources(tmp_path, ImageReader())
    assert isinstance(result, Err)
    assert result.error.code == "invalid_source"


def test_structured_database_preserves_decimal_precision(tmp_path):
    (tmp_path / "precise.json").write_text('{"value":1234567890123.4567}')
    result = load_sources(tmp_path, ImageReader())
    assert isinstance(result, Ok)
    assert result.value.databases[0]["value"] == Decimal("1234567890123.4567")


@pytest.mark.parametrize(
    "complete, changed, approved",
    [(True, False, True), (False, False, False), (True, True, False)],
)
def test_only_complete_stable_image_reply_is_approved_for_cache(
    tmp_path, complete, changed, approved
):
    (tmp_path / "source.png").write_bytes(PNG)

    class ApprovalReader:
        def __init__(self):
            self.approved = []

        def complete(self, **kwargs):
            if changed:
                kwargs["image_path"].write_bytes(b"modified image")
            return Ok(
                ModelReply(
                    "", {"text": "Image evidence", "complete": complete}, {}, "fake"
                )
            )

        def approve_cache(self, reply):
            self.approved.append(reply)
            return True

    reader = ApprovalReader()
    result = load_sources(tmp_path, reader)
    assert isinstance(result, Ok) is approved
    assert bool(reader.approved) is approved


def test_open_word_document_owner_file_is_not_client_evidence(tmp_path):
    doc = Document()
    doc.add_paragraph("Recorded recommendation and reason.")
    doc.save(tmp_path / "report_request.docx")
    reader = ImageReader()
    before = load_sources(tmp_path, reader)
    assert isinstance(before, Ok)
    (tmp_path / "~$port_request.docx").write_bytes(b"Word owner metadata")
    after = load_sources(tmp_path, reader)
    assert isinstance(after, Ok)
    assert after.value == before.value
    assert not reader.calls
    assert (tmp_path / "~$port_request.docx").exists()
